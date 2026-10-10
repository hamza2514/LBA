"""Excel export of the result table (current order, colours, assigned-by)."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

COLUMNS = ["Line", "Operation", "Machine Type", "SAM", "Employee Code", "Employee Name", "Target", "Assigned By"]
_WIDTHS = [14, 52, 20, 9, 15, 26, 10, 13]
_THIN = Side(style="thin", color="D5DAE1")


def to_excel_bytes(records: list[dict], sheet_name: str = "Layout") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    ws.append(COLUMNS)
    header_fill = PatternFill("solid", fgColor="0F172A")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    for rec in records:
        ws.append([rec[c] if rec[c] != "" else None for c in COLUMNS])
        row = ws[ws.max_row]
        fill = PatternFill("solid", fgColor=rec["_color"].lstrip("#")) if rec["_color"] else None
        for cell in row:
            cell.border = Border(top=_THIN, bottom=_THIN, left=_THIN, right=_THIN)
            cell.alignment = Alignment(vertical="center")
            if fill:
                cell.fill = fill
        if not rec["Employee Code"]:
            row[4].value = "UNASSIGNED"
            row[4].font = Font(bold=True, color="B91C1C")

    for i, width in enumerate(_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
